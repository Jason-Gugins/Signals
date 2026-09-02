"""NJ WARN full-archive XLSX parser (openpyxl).

Source: https://www.nj.gov/labor/assets/PDFs/WARN/WARN_Notice_Archive.xlsx
(2004-2026 full archive; column headers verified by P3 live-probe research).
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from src.core.textutil import to_iso_date
from src.sources.warn import WarnNotice, parse_affected


class NjWarn:
    code = "NJ"
    index_url = "https://www.nj.gov/labor/assets/PDFs/WARN/WARN_Notice_Archive.xlsx"
    fmt = "xlsx"

    def discover(self, body: bytes) -> list[str]:
        return []

    def parse(self, body: bytes) -> list[WarnNotice]:
        if not body:
            return []
        wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        out: list[WarnNotice] = []
        header: list[str] | None = None
        for row in rows:
            if header is None:
                header = [str(c).strip().casefold() if c is not None else "" for c in row]
                continue
            rec = {header[i]: ("" if c is None else str(c).strip()) for i, c in enumerate(row) if i < len(header)}
            company = rec.get("company") or rec.get("company name") or ""
            if not company:
                continue  # malformed row: missing company -> skip
            out.append(
                WarnNotice(
                    company_raw=company,
                    state="NJ",
                    notice_date=to_iso_date(rec.get("notice date")),
                    effective_date=to_iso_date(rec.get("effective date")),
                    affected=parse_affected(rec.get("number affected") or rec.get("affected")),
                    location=rec.get("city") or None,
                    url=None,
                    reason=rec.get("reason") or None,
                )
            )
        return out
